"""Load a document package into Document objects.

A package is a directory of mixed files (or a single file). Each file is parsed
by format and tagged with a modality (from its extension via the taxonomy):

  .json         AWS Textract OCR output
  .xlsx/.xlsm   Excel workbook (needs openpyxl)
  .eml          email
  .txt / other  plain text

Raw .pdf must be OCR'd through Textract first (-> .json).
"""

from __future__ import annotations

import json
import logging
import os
import re
from email import message_from_bytes
from email.policy import default as email_policy

from . import taxonomy
from .models import Document

logger = logging.getLogger(__name__)

_EXCEL_MAX_ROWS = 200
_EXCEL_MAX_LINES = 800


# --- per-format parsers (return a dict of fields) ----------------------------

def _parse_textract(path, file_name, doc_id) -> dict:
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    blocks = data.get("Blocks", []) or []
    by_id = {b["Id"]: b for b in blocks if "Id" in b}

    def child_text(block):
        parts = []
        for rel in block.get("Relationships", []) or []:
            if rel.get("Type") == "CHILD":
                for cid in rel.get("Ids", []):
                    c = by_id.get(cid, {})
                    if c.get("BlockType") == "WORD":
                        parts.append(c.get("Text", ""))
        return " ".join(parts).strip()

    lines = [b.get("Text", "") for b in blocks if b.get("BlockType") == "LINE" and b.get("Text")]
    key_values = []
    for b in blocks:
        if b.get("BlockType") == "KEY_VALUE_SET" and "KEY" in (b.get("EntityTypes") or []):
            key = child_text(b)
            value = ""
            for rel in b.get("Relationships", []) or []:
                if rel.get("Type") == "VALUE":
                    for vid in rel.get("Ids", []):
                        value = child_text(by_id.get(vid, {}))
            if key:
                key_values.append(f"{key}={value}".strip())
    pages = (data.get("DocumentMetadata") or {}).get("Pages") or \
        max([b.get("Page", 1) for b in blocks if b.get("BlockType") == "PAGE"] or [1])
    return dict(text="\n".join(lines), lines=lines, key_values=key_values,
                table_count=sum(1 for b in blocks if b.get("BlockType") == "TABLE"),
                page_count=pages)


def _parse_excel(path, file_name, doc_id) -> dict:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Reading Excel needs openpyxl (pip install openpyxl).") from exc
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines, key_values = [], []
    for ws in wb.worksheets:
        lines.append(f"[SHEET] {ws.title}")
        header = False
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            if r >= _EXCEL_MAX_ROWS:
                lines.append("... (truncated)")
                break
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if not cells:
                continue
            if not header:
                header = True
                key_values += [f"column={h}" for h in cells[:20]]
            lines.append(" | ".join(cells))
            if len(lines) >= _EXCEL_MAX_LINES:
                break
        if len(lines) >= _EXCEL_MAX_LINES:
            break
    sheets = list(wb.sheetnames)
    wb.close()
    return dict(text="\n".join(lines), lines=lines, key_values=key_values,
                table_count=len(sheets), page_count=len(sheets))


def _parse_eml(path, file_name, doc_id) -> dict:
    with open(path, "rb") as fh:
        msg = message_from_bytes(fh.read(), policy=email_policy)
    subject = str(msg.get("subject", "")).strip()
    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                body = part.get_content()
                break
    else:
        body = msg.get_content()
    text = "\n".join(p for p in [f"Subject: {subject}" if subject else "", body.strip()] if p)
    kv = [f"{k}={msg.get(k)}" for k in ("subject", "from", "to") if msg.get(k)]
    return dict(text=text, lines=text.splitlines(), key_values=kv,
                table_count=0, page_count=1)


def _parse_text(path, file_name, doc_id) -> dict:
    with open(path, "r", encoding="utf-8", errors="ignore") as fh:
        text = fh.read()
    return dict(text=text, lines=text.splitlines(), key_values=[],
                table_count=0, page_count=1)


_PARSERS = {".json": _parse_textract, ".xlsx": _parse_excel, ".xlsm": _parse_excel,
            ".eml": _parse_eml}


def load_file(path: str, file_name: str, doc_id: str) -> Document:
    ext = os.path.splitext(file_name)[1].lower()
    parser = _PARSERS.get(ext, _parse_text)
    fields = parser(path, file_name, doc_id)
    modality = taxonomy.get().modality_for_extension(ext)
    return Document(doc_id=doc_id, file_name=file_name, modality=modality, **fields)


def load_package(path: str) -> tuple[str, list[Document]]:
    """Return (package_id, [Document, ...]) for a directory or single file."""
    if os.path.isdir(path):
        package_id = os.path.basename(os.path.normpath(path))
        files = sorted(f for f in os.listdir(path)
                       if os.path.isfile(os.path.join(path, f)) and not f.startswith("."))
        if not files:
            raise ValueError(f"No files in package directory: {path}")
        docs = [load_file(os.path.join(path, f), f, f"doc_{i+1}")
                for i, f in enumerate(files)]
    elif os.path.isfile(path):
        package_id = re.sub(r"\.[^.]+$", "", os.path.basename(path))
        docs = [load_file(path, os.path.basename(path), "doc_1")]
    else:
        raise FileNotFoundError(f"Package path not found: {path}")

    logger.info("Loaded package '%s' with %d document(s).", package_id, len(docs))
    return package_id, docs
