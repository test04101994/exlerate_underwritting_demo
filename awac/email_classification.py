"""FNOL email classifier.

Recursively scans for ``.eml`` files, pre-screens them, and uses AWS
Bedrock Claude to classify each as FNOL or Not FNOL. Writes results to
an Excel workbook.

Fill in the AWS credential placeholders below before running.
"""

from __future__ import annotations

import argparse
import email
import json
import logging
import os
import re
import sys
from email.message import Message
from pathlib import Path
from typing import Any

import boto3
import openpyxl
from botocore.exceptions import BotoCoreError, ClientError
from tqdm import tqdm

# ─────────────────────────────────────────────────────────────
# AWS credentials (fill these in before running)
# WARNING: do not commit real credentials to source control.
# ─────────────────────────────────────────────────────────────
AWS_ACCESS_KEY_ID: str = "<YOUR_AWS_ACCESS_KEY_ID>"
AWS_SECRET_ACCESS_KEY: str = "<YOUR_AWS_SECRET_ACCESS_KEY>"
AWS_SESSION_TOKEN: str = "<YOUR_AWS_SESSION_TOKEN>"
AWS_REGION: str = "us-east-1"

# ─────────────────────────────────────────────────────────────
# Configuration
# ─────────────────────────────────────────────────────────────
INFERENCE_PROFILE_ID: str = "<YOUR_APPLICATION_INFERENCE_PROFILE_ID>"
OUTPUT_EXCEL: str = "fnol_classification_results.xlsx"
FNOL_TO_ADDRESS: str = "noticeofloss@awac.com"
ALLIED_DOMAIN: str = "awac.com"

MAX_BODY_CHARS: int = 6000
MAX_OUTPUT_TOKENS: int = 300

EXCLUDED_DIRS: frozenset[str] = frozenset({
    "__pycache__", ".git", ".venv", "venv", "env",
    ".idea", ".vscode", "node_modules",
})

EMAIL_ADDRESS_RE: re.Pattern[str] = re.compile(r"[\w.\-+]+@[\w.\-]+")
JSON_FENCE_RE: re.Pattern[str] = re.compile(r"^```json|```$", re.MULTILINE)

EXCEL_HEADERS: tuple[str, ...] = ("File Path", "Classification", "Confidence", "Reason")

SYSTEM_PROMPT: str = """You are an expert insurance claims classifier working for Allied World Insurer.

You will receive a complete email which may contain a single message or a chain
of replies and forwards. Read the entire content and determine whether anywhere
in this email there is a First Notification of Loss (FNOL) being reported to
Allied World Insurer.

Respond ONLY with valid JSON, no markdown, no extra text:
{
  "is_fnol": true | false,
  "confidence": "High" | "Medium" | "Low",
  "reason": "<2 concise lines max explaining your decision>"
}

Classify as FNOL (true) if anywhere in the email:
- A loss or incident is being reported for the first time
- A date of loss or incident is mentioned
- A policy number is referenced alongside a new claim
- The insured or broker explicitly reports a loss event
- Phrases like 'reporting a loss', 'first notice', 'new claim',
  'incident occurred', 'date of loss' appear

Classify as Not FNOL (false) if the email is:
- An acknowledgement or auto-reply only
- A status update on an existing claim
- A renewal, invoice, or policy document
- Out-of-office or test message
- General inquiry with no loss details"""


logger = logging.getLogger("fnol_classifier")


# ─────────────────────────────────────────────────────────────
# Bedrock client
# ─────────────────────────────────────────────────────────────
def get_bedrock_client() -> Any:
    """Build a Bedrock runtime client from the credential constants above."""
    placeholders = {
        "AWS_ACCESS_KEY_ID": AWS_ACCESS_KEY_ID,
        "AWS_SECRET_ACCESS_KEY": AWS_SECRET_ACCESS_KEY,
        "AWS_SESSION_TOKEN": AWS_SESSION_TOKEN,
        "INFERENCE_PROFILE_ID": INFERENCE_PROFILE_ID,
    }
    unfilled = [name for name, value in placeholders.items()
                if not value or value.startswith("<")]
    if unfilled:
        raise ValueError(
            "Replace the placeholder(s) at the top of this file: "
            + ", ".join(unfilled)
        )

    return boto3.client(
        "bedrock-runtime",
        region_name=AWS_REGION,
        aws_access_key_id=AWS_ACCESS_KEY_ID,
        aws_secret_access_key=AWS_SECRET_ACCESS_KEY,
        aws_session_token=AWS_SESSION_TOKEN,
    )


# ─────────────────────────────────────────────────────────────
# Email discovery
# ─────────────────────────────────────────────────────────────
def discover_emails(root_dir: Path) -> list[Path]:
    """Return sorted ``.eml`` paths under ``root_dir``, skipping hidden/excluded dirs."""
    found: list[Path] = []
    for current_dir, subdirs, files in os.walk(root_dir):
        subdirs[:] = [
            d for d in subdirs
            if d not in EXCLUDED_DIRS and not d.startswith(".")
        ]
        for filename in files:
            if filename.lower().endswith(".eml"):
                found.append(Path(current_dir) / filename)
    return sorted(found)


# ─────────────────────────────────────────────────────────────
# Email parsing
# ─────────────────────────────────────────────────────────────
def get_full_body(msg: Message) -> str:
    """Return concatenated decoded ``text/plain`` body of an email message."""
    parts: list[str] = []

    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() != "text/plain":
                continue
            payload = part.get_payload(decode=True)
            if payload is None:
                continue
            charset = part.get_content_charset() or "utf-8"
            parts.append(payload.decode(charset, errors="replace"))
    else:
        payload = msg.get_payload(decode=True)
        if payload is not None:
            charset = msg.get_content_charset() or "utf-8"
            parts.append(payload.decode(charset, errors="replace"))

    return "".join(parts)


def extract_addresses(text: str) -> list[str]:
    """Return all lowercase email addresses found in ``text``."""
    if not text:
        return []
    return [match.lower().strip() for match in EMAIL_ADDRESS_RE.findall(text)]


def parse_eml(filepath: Path) -> dict[str, Any]:
    """Parse a ``.eml`` file into ``sender_domain``, ``all_recipients``, and ``body``."""
    msg = email.message_from_bytes(filepath.read_bytes())

    to_header = msg.get("To", "") or ""
    cc_header = msg.get("CC", "") or msg.get("Cc", "") or ""
    from_header = msg.get("From", "") or ""

    sender_addrs = extract_addresses(from_header)
    sender_domain = sender_addrs[0].split("@")[-1] if sender_addrs else ""

    return {
        "sender_domain": sender_domain,
        "all_recipients": extract_addresses(f"{to_header} {cc_header}"),
        "body": get_full_body(msg)[:MAX_BODY_CHARS],
    }


# ─────────────────────────────────────────────────────────────
# Pre-screening
# ─────────────────────────────────────────────────────────────
def pre_screen(data: dict[str, Any]) -> tuple[bool, str]:
    """Return ``(passes, reason)`` from cheap routing checks before LLM call."""
    if FNOL_TO_ADDRESS not in data["all_recipients"]:
        return False, f"Not addressed to {FNOL_TO_ADDRESS}"
    if data["sender_domain"] == ALLIED_DOMAIN:
        return False, f"Sender is internal ({data['sender_domain']})"
    return True, ""


# ─────────────────────────────────────────────────────────────
# Bedrock classification
# ─────────────────────────────────────────────────────────────
def _strip_json_fences(text: str) -> str:
    """Strip leading ```json and trailing ``` fences."""
    return JSON_FENCE_RE.sub("", text).strip()


def classify(client: Any, body: str) -> dict[str, Any]:
    """Classify an email body via Bedrock Claude; returns parsed JSON dict."""
    payload = {
        "anthropic_version": "bedrock-2023-05-31",
        "max_tokens": MAX_OUTPUT_TOKENS,
        "system": SYSTEM_PROMPT,
        "messages": [{"role": "user", "content": body}],
    }

    response = client.invoke_model(
        modelId=INFERENCE_PROFILE_ID,
        contentType="application/json",
        accept="application/json",
        body=json.dumps(payload),
    )

    response_body = json.loads(response["body"].read())
    raw_text = response_body["content"][0]["text"].strip()
    cleaned = _strip_json_fences(raw_text)

    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        logger.warning("Failed to parse LLM JSON response: %r", cleaned[:200])
        return {
            "is_fnol": None,
            "confidence": "Low",
            "reason": "LLM response parse error.",
        }


# ─────────────────────────────────────────────────────────────
# Result formatting
# ─────────────────────────────────────────────────────────────
def _label_for(record: dict[str, Any]) -> str:
    """Map classification fields to a human-readable label."""
    if record.get("skipped"):
        return "Skipped"
    is_fnol = record.get("is_fnol")
    if is_fnol is True:
        return "FNOL"
    if is_fnol is False:
        return "Not FNOL"
    return "Unknown"


def write_excel(records: list[dict[str, Any]], output_path: Path) -> None:
    """Write classification ``records`` to an ``.xlsx`` at ``output_path``."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "FNOL Classification"
    sheet.append(list(EXCEL_HEADERS))

    for record in records:
        sheet.append([
            record["file_path"],
            _label_for(record),
            record.get("confidence", ""),
            record.get("reason", ""),
        ])

    workbook.save(str(output_path))


# ─────────────────────────────────────────────────────────────
# Per-file processing
# ─────────────────────────────────────────────────────────────
def process_email(client: Any, filepath: Path, rel_path: str) -> dict[str, Any]:
    """Parse, pre-screen, and classify one email; errors are captured into the record."""
    try:
        data = parse_eml(filepath)
    except Exception as exc:  # noqa: BLE001 - per-file isolation is intentional
        logger.exception("Failed to parse %s", rel_path)
        return {
            "file_path": rel_path,
            "is_fnol": None,
            "confidence": "Low",
            "reason": f"Parse error: {str(exc)[:100]}",
            "skipped": False,
        }

    passes, skip_reason = pre_screen(data)
    if not passes:
        return {
            "file_path": rel_path,
            "is_fnol": None,
            "confidence": "N/A",
            "reason": skip_reason,
            "skipped": True,
        }

    try:
        result = classify(client, data["body"])
    except (BotoCoreError, ClientError, KeyError, ValueError) as exc:
        logger.exception("Classification failed for %s", rel_path)
        result = {
            "is_fnol": None,
            "confidence": "Low",
            "reason": str(exc)[:120],
        }

    return {
        "file_path": rel_path,
        "is_fnol": result.get("is_fnol"),
        "confidence": result.get("confidence", ""),
        "reason": result.get("reason", ""),
        "skipped": False,
    }


# ─────────────────────────────────────────────────────────────
# Test mode (single email smoke test)
# ─────────────────────────────────────────────────────────────
def run_test(script_dir: Path) -> int:
    """Run the full pipeline on the first .eml found and print the result.

    Useful as a smoke test — verifies AWS credentials, inference profile,
    email parsing, and JSON parsing without touching the Excel output.
    """
    client = get_bedrock_client()
    eml_files = discover_emails(script_dir)

    if not eml_files:
        logger.error("No .eml files found under %s.", script_dir)
        return 1

    filepath = eml_files[0]
    rel_path = str(filepath.relative_to(script_dir))
    logger.info("Test: processing %s", rel_path)

    record = process_email(client, filepath, rel_path)
    print(json.dumps(record, indent=2))
    return 0


# ─────────────────────────────────────────────────────────────
# Entry point
# ─────────────────────────────────────────────────────────────
def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--test",
        action="store_true",
        help="Run on a single .eml and print the result (no Excel output).",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    """Run the FNOL classification pipeline; returns process exit code."""
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    )

    args = parse_args(argv)
    script_dir = Path(__file__).parent

    if args.test:
        return run_test(script_dir)

    output_path = script_dir / OUTPUT_EXCEL

    client = get_bedrock_client()
    eml_files = discover_emails(script_dir)

    if not eml_files:
        logger.warning("No .eml files found under %s. Exiting.", script_dir)
        return 1

    logger.info("Found %d .eml file(s) to process.", len(eml_files))

    records: list[dict[str, Any]] = []
    for filepath in tqdm(eml_files, desc="Classifying emails"):
        rel_path = str(filepath.relative_to(script_dir))
        records.append(process_email(client, filepath, rel_path))

    write_excel(records, output_path)
    logger.info("Done. Results saved to %s", output_path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
